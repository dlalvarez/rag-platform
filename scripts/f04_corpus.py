"""Synthetic semantic F04 corpus; no existing F03 idempotency keys are reused."""

import io
import zipfile
from datetime import UTC, datetime

import docx
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject


def pdf(
    pages=("Página uno: SELECT valor FROM tabla;", "Página dos: parámetro técnico."),
    *,
    scanned=False,
    encrypted=False,
):
    writer = PdfWriter()
    for index, text in enumerate(pages):
        page = writer.add_blank_page(width=400, height=300)
        content = DecodedStreamObject()
        if scanned is True or (isinstance(scanned, tuple) and index in scanned):
            image = DecodedStreamObject()
            image.set_data(b"\x20\x80\xc0")
            image.update(
                {
                    NameObject("/Type"): NameObject("/XObject"),
                    NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(1),
                    NameObject("/Height"): NumberObject(1),
                    NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                    NameObject("/BitsPerComponent"): NumberObject(8),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/XObject"): DictionaryObject(
                        {NameObject("/Im1"): writer._add_object(image)}
                    )
                }
            )
            content.set_data(b"q 300 0 0 200 20 20 cm /Im1 Do Q")
        else:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                    NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    )
                }
            )
            content.set_data(
                b"BT /F1 12 Tf 20 250 Td <" + text.encode("cp1252").hex().encode() + b"> Tj ET"
            )
        page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("synthetic-only-not-an-account")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def fixed_zip(body):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(body)) as source, zipfile.ZipFile(output, "w") as target:
        for old in source.infolist():
            info = zipfile.ZipInfo(old.filename, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(old))
    return output.getvalue()


def word():
    document = docx.Document()
    document.core_properties.created = datetime(2000, 1, 1, tzinfo=UTC)
    document.core_properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
    document.add_heading("Operación RAG", level=1)
    document.add_paragraph("SELECT * FROM tabla WHERE clave = 'á';")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Clave", "Valor"
    table.cell(1, 0).text, table.cell(1, 1).text = "réplica", "sí"
    document.add_heading("Recuperación", level=2)
    document.add_paragraph("Conservar el original y su versión.")
    output = io.BytesIO()
    document.save(output)
    return fixed_zip(output.getvalue())


def corpus():
    fence = chr(96) * 3
    markdown = "# Operación\n\nTexto técnico español.\n\n" + fence + "sql\n"
    markdown += "SELECT * FROM tabla WHERE clave = 'á';\n" + fence + "\n\n## Réplicas\nConservar.\n"
    return {
        "pdf": ("application/pdf", pdf()),
        "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", word()),
        "md": ("text/markdown", markdown.encode()),
        "txt": ("text/plain", "Árbol técnico\n  SELECT a, b FROM tabla;\nCafe\u0301.\n".encode()),
    }
